"""Tests for turning CSV and Excel files into databases."""

import hashlib
import io

import pytest
from openpyxl import Workbook

from querypilot.safe_db import SafeDB, UnsafeQueryError
from querypilot.uploads import (UploadError, build_database, clean_name, path_for,
                                suggest_questions, to_value)


def test_french_csv_semicolons_accents_decimal_commas(tmp_path):
    """Check that french csv semicolons accents decimal commas."""
    data = ("Région;Chiffre d'affaires (€);Ventes;Date\n"
            "Île-de-France;1 234,50;12;2025-01-03\n"
            "Bretagne;980,25;7;2025-01-04\n").encode("cp1252")
    info = build_database(data, "ventes 2025.csv", tmp_path)
    t = info["tables"][0]
    assert t["name"] == "ventes_2025"
    assert t["columns"] == ["region", "chiffre_d_affaires", "ventes", "date"]
    assert t["types"] == ["TEXT", "REAL", "INTEGER", "TEXT"]
    db = SafeDB(info["path"])
    assert db.query("SELECT SUM(chiffre_d_affaires) FROM ventes_2025").rows == [(2214.75,)]
    assert db.query("SELECT region FROM ventes_2025 WHERE ventes = 12").rows == [("Île-de-France",)]


def test_xlsx_every_sheet_becomes_a_table(tmp_path):
    """Check that xlsx every sheet becomes a table."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["id", "amount", "country"])
    ws.append([1, 10.5, "FR"])
    ws.append([2, 20, "DE"])
    ws2 = wb.create_sheet("Customers 2025")
    ws2.append(["id", "name"])
    ws2.append([1, "Ana"])
    wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    info = build_database(buf.getvalue(), "shop.xlsx", tmp_path)
    assert [t["name"] for t in info["tables"]] == ["orders", "customers_2025"]
    db = SafeDB(info["path"])
    assert db.query("SELECT SUM(amount) FROM orders").rows == [(30.5,)]


def test_duplicate_blank_and_hostile_column_names(tmp_path):
    """Check that duplicate blank and hostile column names."""
    data = b'a,a,,"x""); DROP TABLE t; --",2nd\n1,2,3,4,5\n'
    info = build_database(data, "t.csv", tmp_path)
    cols = info["tables"][0]["columns"]
    assert cols == ["a", "a_2", "col_3", "x_drop_table_t", "c_2nd"]
    assert SafeDB(info["path"]).query("SELECT COUNT(*) FROM t").rows == [(1,)]


def test_uploaded_db_stays_read_only(tmp_path):
    """Check that uploaded db stays read only."""
    info = build_database(b"x\n1\n", "t.csv", tmp_path)
    before = hashlib.sha256(info["path"].read_bytes()).hexdigest()
    with pytest.raises(Exception):
        SafeDB(info["path"]).query("DELETE FROM t")
    assert hashlib.sha256(info["path"].read_bytes()).hexdigest() == before


@pytest.mark.parametrize("data,name,msg", [
    (b"", "a.csv", "empty"),
    (b"only header\n", "a.csv", "header row"),
    (b"x", "a.pdf", ".csv or .xlsx"),
    (b"not excel", "a.xlsx", "valid .xlsx"),
])
def test_bad_files_get_clear_messages(tmp_path, data, name, msg):
    """Check that bad files get clear messages."""
    with pytest.raises(UploadError, match=msg):
        build_database(data, name, tmp_path)


def test_path_for_rejects_traversal(tmp_path):
    """Check that path for rejects traversal."""
    assert path_for("../../etc/passwd", tmp_path) is None
    assert path_for("0" * 32, tmp_path) is None


def test_values_and_names():
    """Check that values and names."""
    assert to_value("12,5") == 12.5 and to_value(" 1 200 ") == 1200 and to_value("") is None
    assert to_value("0612345678") == 612345678 and to_value("33612345678901234") == "33612345678901234"
    assert clean_name("Chiffre d'affaires (€)", "c") == "chiffre_d_affaires"
    qs = suggest_questions([{"name": "sales", "columns": ["region", "amount"], "types": ["TEXT", "REAL"]}])
    assert "What are the top 5 region by total amount?" in qs
