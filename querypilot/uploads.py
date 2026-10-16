"""Turns an uploaded CSV or Excel file into a small SQLite database."""

from __future__ import annotations

import csv
import io
import re
import sqlite3
import time
import unicodedata
import uuid
from datetime import date, datetime
from pathlib import Path

UPLOAD_DIR = Path("data/uploads")
MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 200_000
MAX_COLS = 100
MAX_TABLES = 10
KEEP_HOURS = 24
DB_ID = re.compile(r"^[0-9a-f]{32}$")


class UploadError(ValueError):
    """Raised when a file can't be turned into a database."""


def clean_name(raw, fallback: str) -> str:
    """Turn any header into a safe SQL name."""
    text = unicodedata.normalize("NFKD", str(raw or "")).encode("ascii", "ignore").decode()
    text = re.sub(r"[^0-9a-zA-Z]+", "_", text).strip("_").lower()
    if not text:
        text = fallback
    if text[0].isdigit():
        text = f"c_{text}"
    return text[:60]


def unique_names(raws: list, prefix: str) -> list[str]:
    """Clean a list of names and make sure none repeat."""
    seen, out = set(), []
    for i, raw in enumerate(raws, 1):
        name = base = clean_name(raw, f"{prefix}_{i}")
        n = 2
        while name in seen:
            name = f"{base}_{n}"
            n += 1
        seen.add(name)
        out.append(name)
    return out


_INT = re.compile(r"^[+-]?\d+$")
_NUM = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")
_COMMA_DEC = re.compile(r"^[+-]?\d+,\d+$")


def to_value(v):
    """Convert one cell to a number, ISO date text, text or empty."""
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, datetime):
        return v.isoformat(sep=" ") if (v.hour or v.minute or v.second) else v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if not s:
        return None
    compact = s.replace(" ", "").replace(" ", "")
    if _INT.match(compact):
        return int(compact) if len(compact.lstrip("+-")) < 16 else s
    if _COMMA_DEC.match(compact):
        return float(compact.replace(",", "."))
    if _NUM.match(compact):
        return float(compact)
    return s


def column_type(values: list) -> str:
    """Pick INTEGER, REAL or TEXT for a column from its values."""
    seen = [v for v in values if v is not None]
    if not seen:
        return "TEXT"
    if all(isinstance(v, int) for v in seen):
        return "INTEGER"
    if all(isinstance(v, (int, float)) for v in seen):
        return "REAL"
    return "TEXT"


def _decode(data: bytes) -> str:
    """Read file bytes as text, trying common encodings."""
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise UploadError("Could not read the file's text encoding.")


def read_csv(data: bytes, name: str) -> list[tuple]:
    """Read a CSV file, detecting its delimiter."""
    text = _decode(data)
    sample = text[:20_000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        class dialect(csv.excel):
            """Fallback CSV format when the delimiter can't be detected."""
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    if len(rows) < 2:
        raise UploadError("The CSV needs a header row and at least one data row.")
    return [(Path(name).stem, rows[0], rows[1:])]


def read_xlsx(data: bytes) -> list[tuple]:
    """Read every non-empty sheet of an Excel file."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as err:
        raise UploadError("This doesn't look like a valid .xlsx file.") from err
    sheets = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)
                if any(c is not None and str(c).strip() for c in r)]
        if len(rows) >= 2:
            sheets.append((ws.title, rows[0], rows[1:]))
    wb.close()
    if not sheets:
        raise UploadError("No sheet has a header row and at least one data row.")
    return sheets


def cleanup(directory: Path = UPLOAD_DIR, hours: float = KEEP_HOURS) -> None:
    """Delete uploaded databases older than 24 hours."""
    if not directory.exists():
        return
    cutoff = time.time() - hours * 3600
    for f in directory.glob("*.sqlite"):
        if f.stat().st_mtime < cutoff:
            f.unlink(missing_ok=True)


def build_database(data: bytes, filename: str, directory: Path = UPLOAD_DIR) -> dict:
    """Convert a CSV or Excel file into a new SQLite database."""
    if not data:
        raise UploadError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise UploadError(f"Files up to {MAX_BYTES // (1024 * 1024)} MB are supported.")
    ext = Path(filename or "").suffix.lower()
    if ext in (".csv", ".tsv", ".txt"):
        sheets = read_csv(data, filename)
    elif ext == ".xlsx":
        sheets = read_xlsx(data)
    else:
        raise UploadError("Please upload a .csv or .xlsx file.")
    sheets = sheets[:MAX_TABLES]

    directory.mkdir(parents=True, exist_ok=True)
    cleanup(directory)
    db_id = uuid.uuid4().hex
    path = directory / f"{db_id}.sqlite"
    con = sqlite3.connect(path)
    tables, truncated = [], False
    try:
        table_names = unique_names([s[0] for s in sheets], "table")
        for tname, (_, header, rows) in zip(table_names, sheets):
            width = min(max(len(header), max((len(r) for r in rows), default=0)), MAX_COLS)
            header = (list(header) + [None] * width)[:width]
            cols = unique_names(header, "col")
            if len(rows) > MAX_ROWS:
                rows, truncated = rows[:MAX_ROWS], True
            values = [[to_value(c) for c in (list(r) + [None] * width)[:width]] for r in rows]
            types = [column_type([row[i] for row in values]) for i in range(width)]
            ddl = ", ".join(f'"{c}" {t}' for c, t in zip(cols, types))
            con.execute(f'CREATE TABLE "{tname}" ({ddl})')
            con.executemany(f'INSERT INTO "{tname}" VALUES ({", ".join("?" * width)})', values)
            tables.append({"name": tname, "columns": cols, "types": types, "rows": len(values)})
        con.commit()
    except Exception:
        con.close()
        path.unlink(missing_ok=True)
        raise
    con.close()
    return {"db_id": db_id, "path": path, "file": Path(filename).name,
            "tables": tables, "truncated": truncated}


def path_for(db_id: str, directory: Path = UPLOAD_DIR) -> Path | None:
    """Find an uploaded database by its id, rejecting invalid ids."""
    if not db_id or not DB_ID.match(db_id):
        return None
    path = directory / f"{db_id}.sqlite"
    return path if path.is_file() else None


def suggest_questions(tables: list[dict]) -> list[str]:
    """Suggest a few starter questions from the uploaded columns."""
    out = []
    for t in tables[:2]:
        nums = [c for c, ty in zip(t["columns"], t["types"]) if ty in ("INTEGER", "REAL")
                and not c.endswith("id")]
        texts = [c for c, ty in zip(t["columns"], t["types"]) if ty == "TEXT"]
        name = t["name"].replace("_", " ")
        out.append(f"How many rows are in {name}?")
        if texts and nums:
            out.append(f"What are the top 5 {texts[0].replace('_', ' ')} by total "
                       f"{nums[0].replace('_', ' ')}?")
        if nums:
            out.append(f"What is the average {nums[0].replace('_', ' ')}?")
        if texts:
            out.append(f"How many distinct {texts[0].replace('_', ' ')} values are there?")
    return out[:5]
